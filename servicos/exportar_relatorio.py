"""servicos/exportar_relatorio.py

Exportação do resumo mensal inteiro (funcionalidade 9 do README): PDF com
fpdf2 e Excel com openpyxl. Módulo dono: Arthur Linhares (Dev 5, seção 3.5).

Diferente de servicos/exportar_pdf.py (que gera o comprovante de UMA
solicitação, funcionalidade 18), este aqui exporta o mês inteiro: total de
horas, dias úteis trabalhados, faltas, saldo acumulado e a evolução dia a
dia - os mesmos números que aparecem em paginas/relatorio.py (Pietro).

Contrato de dados_mes: o MESMO dict que api/rotas/relatorio.py devolve em
GET /relatorio/{estagiario_id} (validado contra RelatorioMensalResponse em
api/schemas.py). Acesso por chave (dict), não por atributo:

    {
        "mes": int, "ano": int,
        "total_horas_trabalhadas": float,
        "dias_uteis_trabalhados": int,
        "faltas": int,
        "saldo_acumulado": float,
        "aviso_vencimento": str | None,
        "evolucao_diaria": [
            {"data": date, "horas_trabalhadas": float, "saldo_do_dia": float},
            ...
        ],
    }

Não recalcula nada - só formata o que api/rotas/relatorio.py já calculou.
Média diária é a única conta feita aqui (não vem pronta do relatorio.py).
"""

from io import BytesIO

from fpdf import FPDF
from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

_NOMES_MES = {
    1: "Janeiro", 2: "Fevereiro", 3: "Março", 4: "Abril", 5: "Maio",
    6: "Junho", 7: "Julho", 8: "Agosto", 9: "Setembro", 10: "Outubro",
    11: "Novembro", 12: "Dezembro",
}


def _media_diaria(dados_mes: dict) -> float:
    """Média diária = total do mês / dias úteis trabalhados.

    Não vem pronta de api/rotas/relatorio.py - é a única conta feita aqui.
    0.0 se ninguém bateu ponto no mês (evita ZeroDivisionError).
    """
    dias = dados_mes["dias_uteis_trabalhados"]
    if dias == 0:
        return 0.0
    return dados_mes["total_horas_trabalhadas"] / dias


# --- PDF (fpdf2) --------------------------------------------------------------


def gerar_relatorio_pdf(dados_mes: dict) -> bytes:
    """PDF do resumo mensal inteiro.

    Devolve bytes prontos pra rota devolver como arquivo
    (Response(content=..., media_type="application/pdf")).
    """
    nome_mes = _NOMES_MES.get(dados_mes["mes"], str(dados_mes["mes"]))
    media = _media_diaria(dados_mes)

    pdf = FPDF(format="A4")
    pdf.add_page()

    pdf.set_font("Helvetica", "B", 16)
    _linha(pdf, "InternHub - Relatório Mensal", altura=10)
    pdf.set_font("Helvetica", "", 11)
    _linha(pdf, f"{nome_mes} de {dados_mes['ano']}")
    pdf.ln(4)

    pdf.set_font("Helvetica", "B", 12)
    _linha(pdf, "Resumo do mês", altura=8)
    pdf.set_font("Helvetica", "", 11)
    _linha(pdf, f"Total de horas trabalhadas: {dados_mes['total_horas_trabalhadas']:.2f}h")
    _linha(pdf, f"Dias úteis trabalhados: {dados_mes['dias_uteis_trabalhados']}")
    _linha(pdf, f"Faltas: {dados_mes['faltas']}")
    _linha(pdf, f"Média diária: {media:.2f}h")
    _linha(pdf, f"Saldo acumulado: {dados_mes['saldo_acumulado']:.2f}h")
    if dados_mes.get("aviso_vencimento"):
        _linha(pdf, f"Aviso: {dados_mes['aviso_vencimento']}")
    pdf.ln(4)

    evolucao = dados_mes.get("evolucao_diaria") or []
    if evolucao:
        pdf.set_font("Helvetica", "B", 12)
        _linha(pdf, "Evolução diária", altura=8)
        _tabela_evolucao(pdf, evolucao)

    return bytes(pdf.output())


def _tabela_evolucao(pdf: FPDF, evolucao: list) -> None:
    """Tabela simples data / horas / saldo do dia, uma linha por dia.

    Sem gerenciar quebra de página na mão: o auto_page_break do fpdf2
    (ligado por padrão) já cuida de pular de página no meio da tabela
    se o mês tiver muitos dias.
    """
    larguras = (40, 60, 60)
    cabecalho = ("Data", "Horas trabalhadas", "Saldo do dia")

    pdf.set_font("Helvetica", "B", 10)
    for largura, texto in zip(larguras, cabecalho):
        pdf.cell(largura, 7, texto, border=1)
    pdf.ln(7)

    pdf.set_font("Helvetica", "", 10)
    for dia in evolucao:
        pdf.cell(larguras[0], 7, dia["data"].strftime("%d/%m/%Y"), border=1)
        pdf.cell(larguras[1], 7, f"{dia['horas_trabalhadas']:.2f}h", border=1)
        pdf.cell(larguras[2], 7, f"{dia['saldo_do_dia']:+.2f}h", border=1)
        pdf.ln(7)


def _linha(pdf: FPDF, texto: str, altura: int = 7) -> None:
    """Escreve uma linha e já pula pra próxima.

    Mesmo helper de servicos/exportar_pdf.py - duplicado de propósito:
    são 3 linhas, não vale a pena criar um módulo utilitário só pra isso.
    """
    pdf.cell(0, altura, texto)
    pdf.ln(altura)


# --- Excel (openpyxl) ----------------------------------------------------------


def gerar_relatorio_excel(dados_mes: dict) -> bytes:
    """Excel do resumo mensal inteiro, com duas abas: Resumo e Evolução diária.

    Devolve bytes prontos pra rota devolver como arquivo (media_type
    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet").
    """
    nome_mes = _NOMES_MES.get(dados_mes["mes"], str(dados_mes["mes"]))
    media = _media_diaria(dados_mes)

    pasta = Workbook()

    resumo = pasta.active
    resumo.title = "Resumo"
    resumo.append(["InternHub - Relatório Mensal"])
    resumo["A1"].font = Font(bold=True, size=14)
    resumo.append([f"{nome_mes} de {dados_mes['ano']}"])
    resumo.append([])
    resumo.append(["Total de horas trabalhadas (h)", round(dados_mes["total_horas_trabalhadas"], 2)])
    resumo.append(["Dias úteis trabalhados", dados_mes["dias_uteis_trabalhados"]])
    resumo.append(["Faltas", dados_mes["faltas"]])
    resumo.append(["Média diária (h)", round(media, 2)])
    resumo.append(["Saldo acumulado (h)", round(dados_mes["saldo_acumulado"], 2)])
    if dados_mes.get("aviso_vencimento"):
        resumo.append(["Aviso", dados_mes["aviso_vencimento"]])

    for linha in resumo.iter_rows(min_row=4, max_row=resumo.max_row, min_col=1, max_col=1):
        linha[0].font = Font(bold=True)
    _ajustar_largura_colunas(resumo)

    evolucao = dados_mes.get("evolucao_diaria") or []
    if evolucao:
        aba_evolucao = pasta.create_sheet("Evolução diária")
        aba_evolucao.append(["Data", "Horas trabalhadas (h)", "Saldo do dia (h)"])
        for celula in aba_evolucao[1]:
            celula.font = Font(bold=True)
        for dia in evolucao:
            aba_evolucao.append([
                dia["data"].strftime("%d/%m/%Y"),
                round(dia["horas_trabalhadas"], 2),
                round(dia["saldo_do_dia"], 2),
            ])
        _ajustar_largura_colunas(aba_evolucao)

    buffer = BytesIO()
    pasta.save(buffer)
    return buffer.getvalue()


def _ajustar_largura_colunas(aba, largura: int = 24) -> None:
    """Largura fixa e generosa nas colunas usadas - evita texto cortado
    sem precisar calcular a largura ideal de cada célula."""
    for indice in range(1, aba.max_column + 1):
        aba.column_dimensions[get_column_letter(indice)].width = largura
